# -*- coding: utf-8 -*-

# Copyright (C) 2026 Libs通用测试软件 作者
# SPDX-License-Identifier: GPL-3.0-or-later
#
# 本软件是自由软件：你可以依据 GNU 通用公共许可证（GPL v3 或更新版本）
# 的条款再分发和/或修改它。详见项目根目录 LICENSE 文件。
#
# 本软件按“现状”分发，不附带任何明示或默示担保；
# 详情请参阅 GNU GPL v3：<https://www.gnu.org/licenses/>。

"""Excel 模板导出单元测试（新功能）。"""
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from libs_analyzer.test_records import TestRecord  # noqa: E402
from libs_analyzer.excel_export import (  # noqa: E402
    export_records_xlsx, templates, template_name,
)


def _sample_records():
    r1 = TestRecord(number=1, timestamp="2026-10-19 10:00",
                    mode="Alloy", spectrum_name="样1", base="steel",
                    operator="张三", n_peaks=77,
                    grades=[{"grade": "316", "score": 96.0}],
                    chemistry={"Fe": {"conc": 66.0, "se": 2.0, "source": "定标"},
                               "Cr": {"conc": 17.2, "se": 0.3, "source": "定标"},
                               "Ni": {"conc": 10.5, "se": 0.4, "source": "定标"}})
    r2 = TestRecord(number=2, timestamp="2026-10-19 11:00",
                    mode="Alloy", spectrum_name="样2", base="steel",
                    grades=[{"grade": "304", "score": 88.0}],
                    chemistry={"Fe": {"conc": 71.0, "se": 2.0, "source": "定标"},
                               "Cr": {"conc": 18.5, "se": 0.3, "source": "定标"},
                               "Ni": {"conc": 8.0, "se": 0.4, "source": "定标"}})
    return [r1, r2]


class TestExcelExport(unittest.TestCase):
    def setUp(self):
        self.path = os.path.join(tempfile.gettempdir(), "libs_test_export.xlsx")
        if os.path.exists(self.path):
            os.remove(self.path)

    def tearDown(self):
        if os.path.exists(self.path):
            os.remove(self.path)

    def _load(self):
        return load_workbook(self.path)

    def test_templates_registry(self):
        self.assertEqual(len(templates()), 3)
        self.assertIn("full", templates())
        self.assertTrue(template_name("grade").startswith("牌号"))

    def test_full_template(self):
        export_records_xlsx(self.path, _sample_records(), template="full")
        wb = self._load()
        self.assertIn("测试记录", wb.sheetnames)
        self.assertIn("化学明细", wb.sheetnames)
        ws = wb["测试记录"]
        # 表头 10 列 + 2 行记录
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.max_column, 10)
        self.assertEqual(ws.cell(row=2, column=1).value, 1)   # 编号
        self.assertEqual(ws.cell(row=2, column=9).value, 96.0)  # 分数
        ws2 = wb["化学明细"]
        self.assertEqual(ws2.max_column, 6)  # 编号/时间/光谱 + Fe/Cr/Ni

    def test_grade_template(self):
        export_records_xlsx(self.path, _sample_records(), template="grade")
        wb = self._load()
        ws = wb["牌号汇总"]
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.max_column, 8)
        self.assertEqual(ws.cell(row=2, column=6).value, "316")

    def test_chemistry_template(self):
        export_records_xlsx(self.path, _sample_records(), template="chemistry")
        wb = self._load()
        ws = wb["化学透视"]
        # 编号/时间/光谱/模式 + Fe/Cr/Ni + CE = 8 列
        self.assertEqual(ws.max_column, 8)
        self.assertEqual(ws.cell(row=2, column=5).value, 66.0)   # Fe
        self.assertEqual(ws.cell(row=2, column=6).value, 17.2)   # Cr
        # CE 值（Fe 基含 C 时才算；此处无 C 为空）
        self.assertIsNone(ws.cell(row=2, column=8).value)


if __name__ == "__main__":
    unittest.main(verbosity=2)
