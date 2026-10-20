# -*- coding: utf-8 -*-

# Copyright (C) 2026 Libs通用测试软件 作者
# SPDX-License-Identifier: GPL-3.0-or-later
#
# 本软件是自由软件：你可以依据 GNU 通用公共许可证（GPL v3 或更新版本）
# 的条款再分发和/或修改它。详见项目根目录 LICENSE 文件。
#
# 本软件按“现状”分发，不附带任何明示或默示担保；
# 详情请参阅 GNU GPL v3：<https://www.gnu.org/licenses/>。

"""
Excel 模板导出（参照 SciAps Z-902 Templates Setup / Excel 导出）。

三种内置模板：
  full       完整报告：测试记录表 + 化学明细（元素×记录 透视）
  grade      牌号汇总：编号/时间/牌号/分数/基体
  chemistry  化学透视：元素×记录 浓度矩阵 + CE 碳当量
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .test_records import calc_ce

_TEMPLATES = {
    "full": "完整报告（记录 + 化学明细）",
    "grade": "牌号汇总",
    "chemistry": "化学透视（元素×记录 + CE）",
}

_HEADER_FILL = PatternFill("solid", fgColor="1F6FB4")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_CENTER = Alignment(horizontal="center")


def templates():
    return list(_TEMPLATES.keys())


def template_name(key):
    return _TEMPLATES.get(key, key)


def _style_header(ws, ncols, nrows=0):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
    if nrows:
        ws.freeze_panes = "A2"
    # 简单列宽
    for c in range(1, ncols + 1):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].width = max(12, min(26, len(str(ws.cell(row=1, column=c).value or "")) + 4))


def _base_info_rows(records):
    rows = []
    for r in records:
        grade = r.grades[0]["grade"] if r.grades else ""
        score = round(r.grades[0]["score"], 1) if r.grades else ""
        ce = calc_ce(r.chemistry)
        rows.append({
            "number": r.number, "timestamp": r.timestamp,
            "mode": "AVG" if r.is_avg else r.mode,
            "spectrum": r.spectrum_name, "base": r.base,
            "operator": r.operator, "peaks": r.n_peaks,
            "grade": grade, "score": score,
            "ce": round(ce["value"], 3) if ce else "",
        })
    return rows


def export_records_xlsx(path, records, template="full"):
    """导出测试记录到 .xlsx。records: list[TestRecord]。"""
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb = Workbook()
    ws = wb.active

    if template == "grade":
        ws.title = "牌号汇总"
        headers = ["编号", "时间", "模式", "光谱", "基体", "牌号", "分数", "操作者"]
        ws.append(headers)
        for b in _base_info_rows(records):
            ws.append([b["number"], b["timestamp"], b["mode"], b["spectrum"],
                       b["base"], b["grade"], b["score"], b["operator"]])
        _style_header(ws, len(headers), len(records))

    elif template == "chemistry":
        ws.title = "化学透视"
        elements = []
        for r in records:
            for el in r.chemistry:
                if el not in elements:
                    elements.append(el)
        headers = ["编号", "时间", "光谱", "模式"] + elements + ["CE(IIW)"]
        ws.append(headers)
        for b, r in zip(_base_info_rows(records), records):
            row = [b["number"], b["timestamp"], b["spectrum"], b["mode"]]
            for el in elements:
                d = r.chemistry.get(el)
                row.append(round(d["conc"], 4) if d and d.get("conc") is not None else "")
            row.append(b["ce"])
            ws.append(row)
        _style_header(ws, len(headers), len(records))

    else:  # full
        ws.title = "测试记录"
        headers = ["编号", "时间", "模式", "光谱", "基体", "操作者", "峰数",
                   "牌号", "分数", "CE(IIW)"]
        ws.append(headers)
        for b in _base_info_rows(records):
            ws.append([b["number"], b["timestamp"], b["mode"], b["spectrum"],
                       b["base"], b["operator"], b["peaks"],
                       b["grade"], b["score"], b["ce"]])
        _style_header(ws, len(headers), len(records))

        # 化学明细 sheet
        ws2 = wb.create_sheet("化学明细")
        elements = []
        for r in records:
            for el in r.chemistry:
                if el not in elements:
                    elements.append(el)
        headers2 = ["编号", "时间", "光谱"] + elements
        ws2.append(headers2)
        for r in records:
            row = [r.number, r.timestamp, r.spectrum_name]
            for el in elements:
                d = r.chemistry.get(el)
                row.append(round(d["conc"], 4) if d and d.get("conc") is not None else "")
            ws2.append(row)
        _style_header(ws2, len(headers2), len(records))

    wb.save(path)
    return path
